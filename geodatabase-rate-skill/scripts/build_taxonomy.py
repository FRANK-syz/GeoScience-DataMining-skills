# -*- coding: utf-8 -*-
"""Build breakpoint/taxonomy_bundle.json from the two GCMD workbooks.
The bundle is the single source of allowed GCMD values (EN + ZH) used by
labeller_compile.py and verify.py.
"""
import argparse, json, os
import openpyxl

def load(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i < 2:                       # two header rows in every sheet
                continue
            if all(x in (None, "") for x in row):
                continue
            rows.append(["" if x is None else str(x) for x in row])
        out[name] = rows
    return out

def find_table(root, name):
    """Locate a GCMD workbook: workspace first, then the skill's bundled assets/."""
    d, cands = os.path.dirname(os.path.abspath(__file__)), [os.path.join(root, name)]
    for _ in range(4):
        cands += [os.path.join(d, name), os.path.join(d, "assets", name)]
        parent = os.path.dirname(d)
        if parent == d:
            break
        d = parent
    for x in cands:
        if os.path.exists(x):
            return os.path.abspath(x)
    return os.path.join(root, name)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=".")
    ap.add_argument("--en", default="地球科学分类.xlsx")
    ap.add_argument("--zh", default="地球科学分类-zh.xlsx")
    a = ap.parse_args()
    root = os.path.abspath(a.root)
    en_path = find_table(root, a.en)
    zh_path = find_table(root, a.zh)
    en = load(en_path)
    zh = load(zh_path)

    def clean(rows):
        return [r for r in rows if not all(x == "" for x in r)]

    bundle = {}
    e, z = clean(en["sciencekeywords"]), clean(zh["sciencekeywords"])
    bundle["sciencekeywords"] = [{"category": x[0], "topic": x[1], "term": x[2],
                                  "topic_zh": y[1], "term_zh": y[2]} for x, y in zip(e, z)]
    e, z = clean(en["platforms"]), clean(zh["platforms"])
    bundle["platforms"] = [{"basis": x[0], "category": x[1],
                            "basis_zh": y[0], "category_zh": y[1]} for x, y in zip(e, z)]
    e, z = clean(en["instruments"]), clean(zh["instruments"])
    bundle["instruments"] = [{"category": x[0], "class": x[1], "type": x[2],
                              "category_zh": y[0], "class_zh": y[1], "type_zh": y[2]} for x, y in zip(e, z)]
    e, z = clean(en["locations"]), clean(zh["locations"])
    bundle["locations"] = [{"category": x[0], "type": x[1],
                            "category_zh": y[0], "type_zh": y[1]} for x, y in zip(e, z)]
    e, z = clean(en["temporalresolutionrange"]), clean(zh["temporalresolutionrange"])
    bundle["temporalresolutionrange"] = [{"en": x[0], "zh": y[0]} for x, y in zip(e, z)]
    e, z = clean(en["horizontalresolutionrange"]), clean(zh["horizontalresolutionrange"])
    bundle["horizontalresolutionrange"] = [{"en": x[0], "zh": y[0]} for x, y in zip(e, z)]

    out = os.path.join(root, "breakpoint", "taxonomy_bundle.json")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    json.dump(bundle, open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("en:", en_path, "| zh:", zh_path)
    for k, v in bundle.items():
        print(k, len(v))
    print("written:", out)

if __name__ == "__main__":
    main()
